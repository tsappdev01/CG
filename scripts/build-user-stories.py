from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY, GOLD, GREY, BAND = "12325C", "D8A62A", "F2F4F7", "FAFBFC"
FONT = "Arial"

# (id, epic, role, story, acceptance criteria, priority, status)
S = [
 # ── Access and identity ───────────────────────────────────────────────────
 ("ACC-01","Access & identity","any user","sign in with my Dubai Investments account rather than a separate CGTOOL password",
  "Given I have a work account, when I choose Sign in with Microsoft, then I am signed in without entering a CGTOOL password; and no CGTOOL password exists for me to set or reset.","Must","Built"),
 ("ACC-02","Access & identity","any user","be told clearly when my login is not yet linked to my person record",
  "Given my sign-in succeeds but no Member record is linked, when I open any screen, then I see a message naming User Management as where an administrator links it, instead of an empty screen.","Must","Built"),
 ("ACC-03","Access & identity","an administrator","link a login to a person record",
  "Given a Member without a login, when I link an Identity account to it, then that person sees their own data at next sign-in; and a login already linked to another Member is refused with a stated reason.","Must","Built"),
 ("ACC-04","Access & identity","the organisation","have idle sessions locked automatically",
  "Given a session has been idle beyond the configured timeout, when I return, then the screen is locked and I must re-authenticate; and work already saved is intact.","Must","Built"),
 ("ACC-05","Access & identity","an administrator","grant each person access per module",
  "Given a Member, when I set the RP & COI, Insider, Related Party Register and RP Transaction access flags, then their menu shows only the granted areas; and every change is written to the audit trail.","Must","Built"),
 ("ACC-06","Access & identity","the organisation","have every user able to sign in via SSO",
  "Given users imported with non-routable (.local) addresses, when migration completes, then each active member has a routable address or a recorded decision that they get no login.","Must","To do"),

 # ── My Register ───────────────────────────────────────────────────────────
 ("REG-01","My Register","any user","record my relatives once and reuse them each period",
  "Given I add a relative with a name and relationship, when a declaration period opens, then that relative is offered on the declaration without re-entry.","Must","Built"),
 ("REG-02","My Register","any user","be stopped from adding the same relative or company twice",
  "Given a matching record already exists, when I save a duplicate, then I am shown the existing record and the save is refused.","Should","Built"),
 ("REG-03","My Register","any user","record companies I or my relatives have an interest in",
  "Given a company, when I record legal name, nature of business and nature of holding — and for a relative's company, which relative — then it appears in my register and on the related declaration step.","Must","Built"),
 ("REG-04","My Register","any user","attach supporting documents to a relative or company",
  "Given a relative or company, when I attach a trade licence or identity document, then it is stored against that record and can be reopened later.","Must","Built"),
 ("REG-05","My Register","a compliance reviewer","see which attached documents have expired",
  "Given a document with an expiry date in the past, when the register is viewed, then it is shown as expired rather than presented as current.","Should","Built"),
 ("REG-06","My Register","a compliance reviewer","trace every change to a register entry",
  "Given any add, edit or delete of a relative or company, when I open the audit log, then I see who made it, what changed from and to, and when.","Must","Built"),

 # ── Insider Declaration ───────────────────────────────────────────────────
 ("INS-01","Insider Declaration","any user","be notified when a declaration period opens",
  "Given a cycle is run for a period, when I am in its recipient list, then I receive an email and the declaration appears on my Declarations screen with its deadline.","Must","Built"),
 ("INS-02","Insider Declaration","any user","declare whether I hold a National Investor Number and DI shares",
  "Given step 1, when I answer the NIN and shareholding questions for myself and my relatives, then the answers are saved against this period's declaration.","Must","Built"),
 ("INS-03","Insider Declaration","any user","have my Emirates ID and passport read automatically on upload",
  "Given I upload an Emirates ID or passport, when it is processed, then the document number, name and expiry date are filled in for me and remain editable before I continue.","Should","Built"),
 ("INS-04","Insider Declaration","any user","save a part-finished declaration and return to it",
  "Given an incomplete declaration, when I save as draft, then my answers are kept; and the declaration still counts as not submitted until I press Submit.","Must","Built"),
 ("INS-05","Insider Declaration","any user","review everything before I submit",
  "Given step 3, when I reach Review, then every answer and uploaded document is shown back to me, and submission is a deliberate separate action.","Must","Built"),
 ("INS-06","Insider Declaration","any user","receive my submitted declaration as a PDF",
  "Given I submit, when the confirmation email is sent, then the completed declaration is attached as a PDF carrying my answers and the submission date.","Must","Built"),
 ("INS-07","Insider Declaration","any user","amend my declaration while the edit window is open",
  "Given the edit window is open, when I reopen and resubmit, then the amendment is recorded with its own timestamp; and once the window closes the screen states that the submission is final.","Must","Built"),

 # ── RP & COI Declaration ──────────────────────────────────────────────────
 ("COI-01","RP & COI Declaration","any user","work through the declaration in named steps",
  "Given the declaration, when I open it, then I see six steps — Relatives, My companies, Relatives' companies, Board roles, Conflicts, Review — and can move between completed steps.","Must","Built"),
 ("COI-02","RP & COI Declaration","a compliance reviewer","distinguish 'nothing to declare' from 'not answered'",
  "Given a section left empty, when I try to continue without ticking 'I have nothing to declare', then I am asked to choose; and the tick is stored as a positive answer.","Must","Built"),
 ("COI-03","RP & COI Declaration","any user","sign my declaration",
  "Given the Review step, when I draw my signature, then it is stored and appears on the declaration and on its PDF beside the date signed.","Must","Built"),
 ("COI-04","RP & COI Declaration","a compliance reviewer","see when someone filed on another's behalf",
  "Given an approved impersonator submits for a member, when the declaration is viewed, then both names are shown — who filed, and who it was filed for.","Must","Built"),
 ("COI-05","RP & COI Declaration","any user","receive my submitted declaration as a PDF",
  "Given I submit, when the confirmation email is sent, then the completed declaration, including my signature, is attached as a PDF.","Must","Built"),

 # ── RP Transactions ───────────────────────────────────────────────────────
 ("RPT-01","RP Transactions","a requestor","raise a transaction against a declared related party",
  "Given Add New Transaction, when I pick a counter-party, then the list is the Related Party Master only; and a party not on it cannot be selected.","Must","Built"),
 ("RPT-02","RP Transactions","a requestor","have my entity and name filled in and locked",
  "Given I open the form, when it loads, then entity and user are pre-filled from my record and cannot be edited.","Must","Built"),
 ("RPT-03","RP Transactions","a requestor","attach the supporting documents",
  "Given a transaction, when I attach a purchase order, contract or quotation, then each is stored against it and visible to the approver and the CCAO.","Must","Built"),
 ("RPT-04","RP Transactions","a requestor","see the pre-checks at submission rather than a silent failure",
  "Given I submit, when the pre-checks run, then I am shown each result — counter-party on the master, value against delegated authority — pass or fail.","Should","Built"),
 ("RPT-05","RP Transactions","the organisation","route a transaction above the delegated authority to a higher approver",
  "Given a value exceeding the approver's limit in the DoA matrix, when it is submitted, then it is routed to the higher authority automatically.","Must","Untested"),
 ("RPT-06","RP Transactions","a requestor","track where my transaction is and who it is with",
  "Given My Transactions, when I open a transaction, then I see its status, who it is currently with by name and job title, and its full history including every amendment.","Must","Built"),
 ("RPT-07","RP Transactions","an approver","see my queue with how long each item has waited",
  "Given Pending My Approval, when I view the queue, then each item shows value, requestor, days waiting and days remaining before escalation.","Must","Built"),
 ("RPT-08","RP Transactions","an approver","return a transaction with remarks",
  "Given a transaction needing change, when I return it with remarks, then the requestor sees the remarks, can amend and resubmit, and the return is kept in the history.","Must","Built"),
 ("RPT-09","RP Transactions","the organisation","escalate an unactioned transaction to the CCAO",
  "Given no approver action for 30 days, when the check runs, then the transaction is escalated to the CCAO and the reason is recorded.","Must","Built"),
 ("RPT-10","RP Transactions","the CCAO","record a governance decision taken offline",
  "Given CCAO Review, when I record the deciding body, meeting date, members who abstained as conflicted and my compliance conclusion, and attach the minutes, then the decision is stored as evidence.","Must","Untested"),
 ("RPT-11","RP Transactions","the CCAO","release an approved transaction to the register",
  "Given a completed governance decision, when I release it, then the transaction appears on the RP Transaction Register with its release date.","Must","Untested"),
 ("RPT-12","RP Transactions","an auditor","be given the complete register of released transactions",
  "Given the register, when I filter and export it, then I receive every released transaction with its approvals, decision and dates.","Must","Built"),
 ("RPT-13","RP Transactions","an administrator","maintain the Delegation of Authority matrix",
  "Given the DoA screen, when I set approval limits per entity and role, then submission routing uses them; and changes are audited.","Must","Untested"),

 # ── Reports ───────────────────────────────────────────────────────────────
 ("REP-01","Reports","a compliance officer","read a report as the page it will be printed as",
  "Given any report, when it opens, then it is shown as an A4 landscape sheet with the Dubai Investments masthead, the headline figures, and Page N of M on every page.","Must","Built"),
 ("REP-02","Reports","a compliance officer","see on every output what the report was filtered to",
  "Given any filter selection, when I print, export to PDF or export to Excel, then the filter statement appears on that output.","Must","Built"),
 ("REP-03","Reports","a compliance officer","export a report to PDF that looks the same for every recipient",
  "Given a report, when I export to PDF, then the file is generated server-side and is identical regardless of the reader's browser or print settings.","Must","Built"),
 ("REP-04","Reports","a compliance officer","export a report to Excel for further work",
  "Given a report, when I export to Excel, then I receive an .xlsx with a frozen, filterable header row and numeric columns stored as numbers.","Must","Built"),
 ("REP-05","Reports","a compliance officer","navigate a long report",
  "Given a multi-page report, when I use the viewer, then I can go first, previous, next and last, type a page number, zoom, fit to width, show all pages and go full screen.","Should","Built"),
 ("REP-06","Reports","a compliance officer","see who has an account and what each can reach",
  "Given the Users report, when I open it, then I see every user with entity, department and each of the four access flags, with counts per flag.","Must","Built"),
 ("REP-07","Reports","a compliance officer","see who has and has not filed an Insider Declaration",
  "Given the Insider Submission report for a period, when I open it, then I see Total Submissions, Non Submissions and Total Users, and one row per notified person with when and by whom it was last modified.","Must","Built"),
 ("REP-08","Reports","a compliance officer","see the same for the RP & COI Declaration",
  "Given the RP & COI Submission report for a period, then it shows the same figures and columns as the Insider one.","Must","To do"),
 ("REP-09","Reports","a compliance officer","read a single person's declaration in full",
  "Given the detail reports, when I filter by name, then I see that person's declared NINs, shareholdings, related parties and conflicts on one sheet.","Must","To do"),
 ("REP-10","Reports","a compliance officer","report share dealing over a date range",
  "Given the Insider Trading report, when I set start, end and file dates, then I see NIN, shareholder, type, date, value and volume for the period.","Should","To do"),

 # ── Notifications ─────────────────────────────────────────────────────────
 ("NOT-01","Notifications","an administrator","schedule a declaration cycle for a period",
  "Given Declaration Setup, when I configure a cycle with its period and recipients, then notifications are sent to exactly that list and the period is what the reports group by.","Must","Built"),
 ("NOT-02","Notifications","an administrator","see exactly who was sent what and when",
  "Given Pending Notifications, when a user says they received nothing, then I can show the send record for that person and cycle.","Must","Built"),
 ("NOT-03","Notifications","any user","be reminded before a deadline",
  "Given an open period with my declaration outstanding, when a reminder is due, then I receive it; and the reminder stops once I submit.","Should","Built"),

 # ── Migration ─────────────────────────────────────────────────────────────
 ("MIG-01","Migration","the organisation","move every legacy user into CGTOOL without loss",
  "Given the legacy database, when the migration runs, then every user is present in CGTOOL or listed with a stated reason for exclusion; and credentials are not carried across.","Must","Partly built"),
 ("MIG-02","Migration","the organisation","run the migration safely and repeatably",
  "Given the migration, when run with the dry-run default, then it reports what it would do and writes nothing; and running it twice creates no duplicates and overwrites nothing edited since.","Must","Partly built"),
 ("MIG-03","Migration","a compliance reviewer","sign off the migration against a reconciliation report",
  "Given a completed run, when I open the reconciliation report, then I see counts on both sides, every record not migrated with the reason, and every decision left to a person.","Must","Partly built"),
 ("MIG-04","Migration","the organisation","move the historic declarations and transactions",
  "Given legacy declarations, related party master records and transactions, when the migration runs, then each is present in CGTOOL with its original dates and its attached documents re-linked.","Must","To do"),
 ("MIG-05","Migration","an administrator","add users from a spreadsheet",
  "Given the user import template, when I load it and run the import, then valid rows are created and every rejected row is reported with the reason and its sheet row number.","Should","Built"),
]

wb = Workbook()
ws = wb.active
ws.title = "User stories"

ws["A1"] = "CGTOOL — User story backlog"
ws["A1"].font = Font(name=FONT, size=15, bold=True, color=NAVY)
ws["A2"] = ("Derived from the CGTOOL User Guide. Status says what exists today: Built = working in UAT; "
            "Untested = written but not yet proven against a real database; Partly built = some of it exists; "
            "To do = not started. Estimate the Untested, Partly built and To do rows.")
ws["A2"].font = Font(name=FONT, size=10, italic=True, color="56606B")
ws.merge_cells("A1:H1"); ws.merge_cells("A2:H2")
ws.row_dimensions[2].height = 30

HEAD = [("ID",10),("Epic",20),("As a…",22),("I want to…",52),("Acceptance criteria",74),("Priority",11),("Status",14),("Estimate (days)",15)]
thin = Side(style="thin", color="D3DCE8")
for i,(h,w) in enumerate(HEAD, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
    c = ws.cell(row=4, column=i, value=h)
    c.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(wrap_text=True, vertical="center")
    c.border = Border(bottom=Side(style="medium", color=GOLD))

for r,(sid,epic,role,story,ac,pri,status) in enumerate(S, start=5):
    vals = [sid, epic, role, story, ac, pri, status, ""]
    for i,v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = Font(name=FONT, size=10, bold=(i==1))
        c.alignment = Alignment(wrap_text=(i in (4,5)), vertical="top")
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        if r % 2 == 0:
            c.fill = PatternFill("solid", fgColor=BAND)
    # Status colour, so the rows that need estimating stand out.
    sc = ws.cell(row=r, column=7)
    tint = {"Built":"E9F8EF","Untested":"FEF3E2","Partly built":"FEF3E2","To do":"FDECEC"}[status]
    sc.fill = PatternFill("solid", fgColor=tint)
    sc.font = Font(name=FONT, size=10, bold=True,
                   color={"Built":"16A34A","Untested":"B45309","Partly built":"B45309","To do":"DC2626"}[status])
    ws.cell(row=r, column=8).fill = PatternFill("solid", fgColor="FFF6D8")

ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:H{4+len(S)}"

# ── Summary sheet (counts typed as values; no formulas, deliberately) ──
sm = wb.create_sheet("Summary")
sm["A1"] = "Backlog summary"
sm["A1"].font = Font(name=FONT, size=13, bold=True, color=NAVY)
sm["A2"] = "Counts are values, not formulas, so this file needs no recalculation to be read correctly."
sm["A2"].font = Font(name=FONT, size=10, italic=True, color="56606B")
sm.merge_cells("A2:D2")

from collections import Counter
by_epic = Counter(x[1] for x in S)
by_status = Counter(x[6] for x in S)
by_pri = Counter(x[5] for x in S)

def block(col, title, counter):
    sm.column_dimensions[get_column_letter(col)].width = 26
    sm.column_dimensions[get_column_letter(col+1)].width = 10
    h = sm.cell(row=4, column=col, value=title)
    h.font = Font(name=FONT, size=9, bold=True, color="FFFFFF"); h.fill = PatternFill("solid", fgColor=NAVY)
    h2 = sm.cell(row=4, column=col+1, value="Stories")
    h2.font = Font(name=FONT, size=9, bold=True, color="FFFFFF"); h2.fill = PatternFill("solid", fgColor=NAVY)
    for r,(k,v) in enumerate(sorted(counter.items(), key=lambda kv: -kv[1]), start=5):
        sm.cell(row=r, column=col, value=k).font = Font(name=FONT, size=10)
        sm.cell(row=r, column=col+1, value=v).font = Font(name=FONT, size=10)
        sm.cell(row=r, column=col).fill = PatternFill("solid", fgColor=GREY)
        sm.cell(row=r, column=col+1).fill = PatternFill("solid", fgColor=GREY)

block(1, "Epic", by_epic)
block(4, "Status", by_status)
block(7, "Priority", by_pri)
sm.cell(row=4+len(by_epic)+2, column=1, value=f"Total stories: {len(S)}").font = Font(name=FONT, size=11, bold=True, color=NAVY)

wb.save("/home/user/cg/docs/CGTOOL-User-Stories.xlsx")
print("stories:", len(S), "| epics:", len(by_epic), "| status:", dict(by_status))
